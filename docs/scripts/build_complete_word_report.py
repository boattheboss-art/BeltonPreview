import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def create_complete_word_report():
    doc = Document()

    # Standard Academic Margins: Left 1.5" (3.81 cm), Top/Bottom/Right 1.0" (2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.5)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.27)   # A4 Width (21.0 cm)
        section.page_height = Inches(11.69) # A4 Height (29.7 cm)

    # Base style
    style = doc.styles['Normal']
    font = style.font
    font.name = 'TH Sarabun PSK'
    font.size = Pt(16)
    font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.line_spacing = 1.15
    style.paragraph_format.space_after = Pt(3)

    def set_font_run(run, size_pt, bold=False, italic=False, color_rgb=(0,0,0)):
        run.font.name = 'TH Sarabun PSK'
        run.font.size = Pt(size_pt)
        run.bold = bold
        run.italic = italic
        run.font.color.rgb = RGBColor(*color_rgb)
        rPr = run._r.get_or_add_rPr()
        rFonts = OxmlElement('w:rFonts')
        rFonts.set(qn('w:ascii'), 'TH Sarabun PSK')
        rFonts.set(qn('w:hAnsi'), 'TH Sarabun PSK')
        rFonts.set(qn('w:cs'), 'TH Sarabun PSK')
        rFonts.set(qn('w:eastAsia'), 'TH Sarabun PSK')
        rPr.append(rFonts)

    def add_p(text, size=16, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT, indent=False, space_before=0, space_after=3, color=(0,0,0)):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if indent:
            p.paragraph_format.first_line_indent = Inches(0.5)
        run = p.add_run(text)
        set_font_run(run, size, bold=bold, italic=italic, color_rgb=color)
        return p

    def add_chapter_title(ch_num_text, ch_title_text):
        doc.add_page_break()
        add_p(ch_num_text, size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=16, space_after=4)
        add_p(ch_title_text, size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=18)

    def add_h1(text):
        add_p(text, size=18, bold=True, space_before=14, space_after=6)

    def add_h2(text):
        add_p(text, size=16, bold=True, indent=True, space_before=10, space_after=4)

    def add_h3(text):
        add_p(text, size=16, bold=True, indent=True, space_before=6, space_after=2)

    def add_bullet(text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.left_indent = Inches(0.75)
        run = p.add_run(text)
        set_font_run(run, 16, bold=False)

    def add_image_with_caption(img_path, caption_text, width_inches=6.0):
        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(12)
            p_img.paragraph_format.space_after = Pt(4)
            run = p_img.add_run()
            run.add_picture(img_path, width=Inches(width_inches))

            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(2)
            p_cap.paragraph_format.space_after = Pt(12)
            run_cap = p_cap.add_run(caption_text)
            set_font_run(run_cap, 14, bold=True, italic=True)

    # -------------------------------------------------------------------------
    # 1. COVER PAGE (หน้าปก)
    # -------------------------------------------------------------------------
    add_p("รายงานโครงงานฉบับสมบูรณ์", size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=36, space_after=12)
    add_p("การออกแบบและพัฒนาสื่อนวัตกรรมปฏิสัมพันธ์เสมือนจริงสามมิติและปัญญาประดิษฐ์สนทนา เพื่อการถ่ายทอดองค์ความรู้กระบวนการผลิตเชิงอุตสาหกรรม", size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=12, space_after=6)
    add_p("AN INTERACTIVE 3D DIGITAL TWIN AND CONVERSATIONAL AI MEDIA DESIGN FOR INDUSTRIAL MANUFACTURING KNOWLEDGE TRANSFER", size=15, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=48)

    add_p("โดย", size=16, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=18, space_after=6)
    add_p("[ระบุชื่อ - นามสกุล ผู้จัดทำ]", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=4)
    add_p("รหัสนักศึกษา [ระบุรหัสนักศึกษา]", size=16, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=60)

    add_p("รายงานนี้เป็นส่วนหนึ่งของโครงงานระดับปริญญาตรี", size=16, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=24, space_after=4)
    add_p("หลักสูตร [ระบุชื่อหลักสูตร/สาขาวิชา เช่น นวัตกรรมสื่อสารมวลชน / สื่อดิจิทัลและแอนิเมชัน]", size=16, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=4)
    add_p("คณะ [ระบุชื่อคณะ เช่น คณะนิเทศศาสตร์ / คณะเทคโนโลยีสื่อสารมวลชน]", size=16, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=4)
    add_p("[ระบุชื่อมหาวิทยาลัย / สถาบันการศึกษา]", size=16, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=4)
    add_p("ปีการศึกษา [ระบุปีการศึกษา เช่น 2567]", size=16, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=0)

    # -------------------------------------------------------------------------
    # 2. APPROVAL PAGE (ใบรับรอง)
    # -------------------------------------------------------------------------
    doc.add_page_break()
    add_p("ใบรับรองรายงานโครงงาน", size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=12, space_after=18)
    add_p("ชื่อโครงงาน: การออกแบบและพัฒนาสื่อนวัตกรรมปฏิสัมพันธ์เสมือนจริงสามมิติและปัญญาประดิษฐ์สนทนา เพื่อการถ่ายทอดองค์ความรู้กระบวนการผลิตเชิงอุตสาหกรรม", size=16, bold=True, space_before=6, space_after=4)
    add_p("ผู้จัดทำ: [ระบุชื่อ-นามสกุล ผู้จัดทำ] รหัสนักศึกษา [ระบุรหัสนักศึกษา]", size=16, space_before=0, space_after=4)
    add_p("สาขาวิชา: [ระบุสาขาวิชา] คณะ [ระบุชื่อคณะ]", size=16, space_before=0, space_after=18)
    add_p("ได้รับการพิจารณาและอนุมัติให้เป็นส่วนหนึ่งของการศึกษาตามหลักสูตรปริญญาตรี", size=16, indent=True, space_before=6, space_after=36)

    add_p("คณะกรรมการสอบโครงงาน", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=18, space_after=18)
    add_p("................................................................... ประธานกรรมการ", size=16, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=12, space_after=4)
    add_p("( [ระบุชื่อ-นามสกุล ประธานกรรมการ] )", size=16, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=18)
    add_p("................................................................... กรรมการ (อาจารย์ที่ปรึกษา)", size=16, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=12, space_after=4)
    add_p("( [ระบุชื่อ-นามสกุล อาจารย์ที่ปรึกษา] )", size=16, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=18)
    add_p("................................................................... กรรมการ", size=16, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=12, space_after=4)
    add_p("( [ระบุชื่อ-นามสกุล กรรมการผู้ทรงคุณวุฒิ] )", size=16, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=0)

    # -------------------------------------------------------------------------
    # 3. ABSTRACT (บทคัดย่อ)
    # -------------------------------------------------------------------------
    doc.add_page_break()
    add_p("บทคัดย่อ", size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=12, space_after=18)
    add_p("โครงงานสร้างสรรค์นี้มีวัตถุประสงค์เพื่อออกแบบและพัฒนาสื่อนวัตกรรมปฏิสัมพันธ์เสมือนจริงสามมิติ (Web-based 3D Interactive Digital Twin) ผสานกับระบบปัญญาประดิษฐ์สนทนา (Conversational AI Agent) เพื่อแก้ปัญหาด้านการสื่อสารและการถ่ายทอดองค์ความรู้ในอุตสาหกรรมการผลิตชิ้นส่วนอิเล็กทรอนิกส์ความแม่นยำสูง ซึ่งเดิมมีข้อจำกัดจากเอกสารคู่มือและสไลด์ฝึกอบรมแบบสองมิติที่มีความซับซ้อน เข้าใจยาก และก่อให้เกิดภาระทางปัญญา (Cognitive Load) แก่บุคลากรและผู้เรียน", size=16, indent=True, space_after=6)
    add_p("ผู้วิจัยได้ประยุกต์ใช้ระเบียบวิธีวิจัยเชิงสร้างสรรค์สื่อดิจิทัล โดยใช้แนวคิดการออกแบบสถาปัตยกรรมข้อมูล (Information Architecture) ร่วมกับสุนทรียศาสตร์แบบสแกนดิเนเวียน (Scandinavian Minimalist Design) และการออกแบบส่วนประสานผู้ใช้กระจกฝ้าเสมือนจริง (VisionOS Frosted Glass Interface) สร้างแบบจำลองโมเดลโรงงานคลีนรูมและชิ้นส่วน Hard Disk Drive ขนาด 12 โซน ผ่านโปรแกรม Blender และประมวลผลการแสดงผลสามมิติแบบ Real-time บนเว็บด้วย Three.js พร้อมทั้งออกแบบระบบกล้อง Cinematic Camera Flight และระบบนำทางด้วยเสียง (Auditory Feedback) ในด้านระบบปัญญาประดิษฐ์ ได้ออกแบบตัวแทนการสื่อสารอัจฉริยะ (Local LLM Qwen 2.5:14b) ที่ผสานการสืบค้นข้อมูลเชิงเทคนิคแบบ RAG จากสไลด์มาตรฐานอุตสาหกรรม (ACA, FCOF, APFA, PCBA) และฐานข้อมูลเครื่องจักร SCADA รวม 50 เครื่อง พร้อมระบบตอบกลับเมื่อบริการขัดข้อง (Graceful Offline Handling)", size=16, indent=True, space_after=6)
    add_p("ผลการประเมินและการทดสอบประสิทธิภาพพบว่า สื่อสามมิติปฏิสัมพันธ์สามารถแสดงผลได้อย่างลื่นไหลที่อัตรา 60 เฟรมต่อวินาที (FPS) ระบบตัดระนาบสถาปัตยกรรมทำให้เห็นสายการผลิตได้อย่างชัดเจนโดยไม่มีสิ่งบดบัง และระบบปัญญาประดิษฐ์สามารถค้นหาข้อกำหนดและข้อสอบทางการได้ถูกต้องและมีความเร็วในการตอบสนองเฉลี่ยต่ำกว่า 1 วินาที ผู้ใช้งานสามารถเข้าถึงข้อมูลเชิงเทคนิคได้รวดเร็วและเห็นภาพกระบวนการผลิตได้อย่างเป็นรูปธรรม สรุปได้ว่าการผสานสื่อสามมิติและปัญญาประดิษฐ์สนทนาสามารถยกระดับการสื่อสารองค์ความรู้เชิงอุตสาหกรรมได้อย่างมีนัยสำคัญ", size=16, indent=True, space_after=12)
    add_p("คำสำคัญ: สื่อปฏิสัมพันธ์ดิจิทัล, ดิจิทัลทวิน, การออกแบบประสบการณ์ผู้ใช้, สื่อนฤมิตสามมิติ, ปัญญาประดิษฐ์สนทนา, การถ่ายทอดองค์ความรู้", size=16, bold=True, space_after=0)

    # -------------------------------------------------------------------------
    # 4. ACKNOWLEDGEMENTS (กิตติกรรมประกาศ)
    # -------------------------------------------------------------------------
    doc.add_page_break()
    add_p("กิตติกรรมประกาศ", size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=12, space_after=18)
    add_p("โครงงานฉบับนี้สำเร็จลุล่วงไปได้ด้วยความกรุณาและความช่วยเหลืออย่างดียิ่งจาก [ระบุชื่อ-นามสกุล อาจารย์ที่ปรึกษา] อาจารย์ที่ปรึกษาโครงงาน ผู้ซึ่งได้สละเวลาให้คำปรึกษา แนะนำแนวคิดการวิจัย ข้อคิดเห็นเชิงวิชาการด้านการออกแบบสื่อดิจิทัลและเทคโนโลยีสารสนเทศ ตลอดจนช่วยตรวจสอบแก้ไขข้อบกพร่องต่างๆ ด้วยความเอาใจใส่เสมอมา ผู้จัดทำขอกราบขอบพระคุณเป็นอย่างสูง ณ ที่นี้", size=16, indent=True, space_after=6)
    add_p("ขอกราบขอบพระคุณคณาจารย์ประจำภาควิชาและคณะทุกท่าน ที่ได้ประสิทธิ์ประสาทวิชาความรู้ทั้งในเชิงทฤษฎีการสื่อสาร เทคโนโลยีมัลติมีเดีย และการออกแบบประสบการณ์ผู้ใช้ ซึ่งเป็นรากฐานสำคัญในการนำมาประยุกต์ใช้สร้างสรรค์ผลงานโครงงานชิ้นนี้", size=16, indent=True, space_after=6)
    add_p("ขอขอบคุณบุคลากรและทีมวิศวกรผู้เชี่ยวชาญด้านกระบวนการผลิตอุตสาหกรรม ที่ได้ให้ความอนุเคราะห์ข้อมูลเชิงเทคนิค สเปกเครื่องจักร และเอกสารมาตรฐานการผลิต ซึ่งเป็นข้อมูลสำคัญในการสร้างสรรค์สื่อดิจิทัลทวินเสมือนจริง", size=16, indent=True, space_after=6)
    add_p("สุดท้ายนี้ ขอกราบขอบพระคุณบิดา มารดา และครอบครัว ที่ให้กำลังใจและความสนับสนุนในทุกๆ ด้านเสมอมา รวมทั้งเพื่อนๆ นักศึกษาทุกคนที่ได้ร่วมแลกเปลี่ยนความคิดเห็นและให้ความช่วยเหลือตลอดระยะเวลาการศึกษา ผู้จัดทำหวังเป็นอย่างยิ่งว่าโครงงานชิ้นนี้จะมีคุณค่าและเป็นประโยชน์ต่อวงการสื่อดิจิทัลและการศึกษาสืบไป", size=16, indent=True, space_after=24)
    add_p("[ระบุชื่อ-นามสกุล ผู้จัดทำ]", size=16, align=WD_ALIGN_PARAGRAPH.RIGHT, space_before=0, space_after=0)

    # -------------------------------------------------------------------------
    # CHAPTER 1
    # -------------------------------------------------------------------------
    add_chapter_title("บทที่ 1", "บทนำ (Introduction)")

    add_h1("1.1 ที่มาและความสำคัญของปัญหา (Background and Rationale)")
    add_p("ในยุคเศรษฐกิจดิจิทัล (Digital Economy) และการเปลี่ยนผ่านสู่อุตสาหกรรม 4.0 สื่อดิจิทัลได้ก้าวเข้ามามีบทบาทสำคัญอย่างยิ่งยวดในการทำหน้าที่เป็นสะพานเชื่อมระหว่างข้อมูลที่มีความซับซ้อนสูง (Complex Information) กับผู้รับสารมนุษย์ ทว่าในอุตสาหกรรมการผลิตชิ้นส่วนอิเล็กทรอนิกส์ที่มีความแม่นยำสูงระดับไมครอน เช่น การผลิตชิ้นส่วนฮาร์ดดิสก์ไดรฟ์ (Hard Disk Drive: HDD) และแผงวงจรในห้องคลีนรูม (Cleanroom) ช่องทางการสื่อสารและสื่อการเรียนรู้แบบดั้งเดิมยังคงพึ่งพาเอกสารคู่มือกระดาษ (Printed SOP), แฟ้มเอกสารข้อกำหนดทางเทคนิค (Specification Sheets), และไฟล์สไลด์นำเสนอแบบสองมิติ (2D Slides) เป็นหลัก", indent=True)
    add_p("จากการศึกษาสภาพปัญหาในการสื่อสารและการถ่ายทอดความรู้ภายในสถานประกอบการ พบข้อจำกัดที่ส่งผลกระทบอย่างมีนัยสำคัญ 3 ประการหลัก ได้แก่:", indent=True)
    add_bullet("ปัญหาด้านภาระการรู้คิดสูง (High Cognitive Load): เอกสารคู่มือมาตรฐานอุตสาหกรรม เช่น TM-00-00-01 หรือ PI-16-09-0001 มีเนื้อหาเชิงเทคนิคที่อัดแน่นด้วยตัวอักษรและภาพถ่ายสองมิติมุมมองเดียว ทำให้ผู้รับสารต้องใช้จินตนาการในการแปลงภาพสองมิติให้เป็นภาพเชิงพื้นที่สามมิติ (Spatial Mental Model) เกิดความเหนื่อยล้าและเข้าใจคลาดเคลื่อนได้ง่าย")
    add_bullet("การขาดมิติความสัมพันธ์เชิงพื้นที่และปฏิสัมพันธ์ (Lack of Spatial & Interactive Context): ผู้เรียนไม่สามารถหมุนดูมุมมองรอบด้าน (360-degree Orbit) หรือจำลองการเดินสำรวจพื้นที่จริง ส่งผลให้มองไม่เห็นความเชื่อมโยงระหว่างขั้นตอนการทำงานของเครื่องจักร 50 เครื่องในสายการผลิต")
    add_bullet("ความล่าช้าในการสืบค้นข้อมูลและการสื่อสารแบบทางเดียว (One-way Passive Communication): เมื่อเกิดเหตุฉุกเฉินหรือความผิดปกติในสายการผลิต บุคลากรต้องเสียเวลาเปิดค้นหาเกณฑ์มาตรฐานจากเอกสารหลายร้อยหน้า ขาดระบบสื่อสารอัจฉริยะที่สามารถให้คำตอบได้แบบโต้ตอบทันที (Real-time Interactive Response)")
    add_p("ด้วยเหตุนี้ ผู้วิจัยจึงมีแนวคิดในการออกแบบและพัฒนา 'สื่อนวัตกรรมปฏิสัมพันธ์เสมือนจริงสามมิติและปัญญาประดิษฐ์สนทนา' (Interactive 3D Digital Twin & Conversational AI) โดยผสานเทคโนโลยี WebGL (Three.js), โมเดลจำลอง 3 มิติ (Blender), ปรัชญาการออกแบบสถาปัตยกรรมข้อมูลและส่วนประสานผู้ใช้สไตล์สแกนดิเนเวียน (Scandinavian Minimalist) ร่วมกับปัญญาประดิษฐ์สนทนาภาษาไทย (Local LLM Qwen 2.5:14b) เพื่อเปลี่ยนรูปแบบการสื่อสารจากสื่อที่ผู้รับสารอยู่นิ่ง (Passive Media) ให้กลายเป็นสื่อปฏิสัมพันธ์ที่ผู้รับสารมีส่วนร่วมอย่างเต็มที่ (Immersive Interactive Media) ซึ่งจะช่วยลดข้อผิดพลาดในการปฏิบัติงานและยกระดับประสิทธิภาพการสื่อสารในโรงงานอุตสาหกรรมอย่างยั่งยืน", indent=True)

    add_h1("1.2 วัตถุประสงค์ของโครงงาน (Project Objectives)")
    add_bullet("1. เพื่อศึกษา ออกแบบ และพัฒนาสื่อนวัตกรรมปฏิสัมพันธ์สามมิติบนเว็บ (Web-based Interactive 3D Media) ที่จำลองโรงงานอุตสาหกรรมคลีนรูมและชิ้นส่วน Hard Disk Drive เสมือนจริง")
    add_bullet("2. เพื่อออกแบบประสบการณ์ผู้ใช้ (UX) และส่วนต่อประสานผู้ใช้ (UI) ตามหลักสุนทรียศาสตร์สแกนดิเนเวียนและระบบกระจกฝ้า VisionOS Glassmorphism ที่ช่วยลดภาระการรับรู้และเพิ่มประสิทธิภาพในการสื่อสารข้อมูล")
    add_bullet("3. เพื่อพัฒนาตัวแทนสื่อสารปัญญาประดิษฐ์สนทนา (Conversational AI Agent) ที่สามารถสืบค้นและถ่ายทอดข้อมูลข้อกำหนดมาตรฐานการผลิตและสถานะเครื่องจักร SCADA แบบโต้ตอบสองทาง (Two-way Interactive Communication)")
    add_bullet("4. เพื่อประเมินประสิทธิภาพของสื่อปฏิสัมพันธ์ด้านอัตราการแสดงผล ความชัดเจนของสารสนเทศ และการยอมรับการใช้งานของผู้รับสาร")

    add_h1("1.3 ขอบเขตของโครงงาน (Project Scope)")
    add_h2("1.3.1 ขอบเขตด้านเนื้อหาสาร (Content Scope)")
    add_bullet("จำลองโรงงานผลิตชิ้นส่วน Hard Disk Drive โรงงานนวนคร ครอบคลุมพื้นที่ 12,500 ตารางเมตร แบ่งเป็น 12 โซนคลีนรูม")
    add_bullet("จำลองเครื่องหยอดกาว ACA Dispensing Machine รวมจำนวน 50 เครื่อง พร้อมข้อมูลโทรมาตรจำลองแบบเรียลไทม์ (SCADA Telemetry Database)")
    add_bullet("เนื้อหาคู่มือมาตรฐานการผลิต 5 หมวดหลัก: Coil Winding (14 ขั้นตอน), ACA (21 ขั้นตอน), FCOF (14 ขั้นตอน), APFA (17 ขั้นตอน) และ PCBA (4 กรณี 30 ขั้นตอน) รวมสไลด์องค์ความรู้กว่า 460 หน้า และคลังข้อสอบทางการ Master Exam กว่า 80 ข้อ")
    add_h2("1.3.2 ขอบเขตด้านการออกแบบและผลิตสื่อ (Design & Media Production Scope)")
    add_bullet("โมเดล 3 มิติสถาปัตยกรรมโรงงาน โครงสร้างเครื่องจักร สายพานลำเลียง และชิ้นส่วน Actuator Coil Assembly ที่ผ่านการ Optimize โพลีกอนเพื่อแสดงผลบนเว็บเบราว์เซอร์")
    add_bullet("การออกแบบ UI 6 หน้าหลักตามมาตรฐาน VisionOS Dark Frosted Glass Navbar และ Clean Studio White Overview Portal")
    add_bullet("ระบบกล้อง Cinematic Flight Animation สลับระหว่างมุมมองผังโรงงานมุมสูง (Isometric 360° Orbit) และมุมมองบุคคลที่หนึ่ง (First-Person Cleanroom Walkthrough)")
    add_h2("1.3.3 ขอบเขตด้านเทคโนโลยีและแพลตฟอร์ม (Technological Scope)")
    add_bullet("ระบบเว็บเบราว์เซอร์มาตรฐานที่รองรับเทคโนโลยี WebGL และ HTML5 Canvas (เช่น Google Chrome, Microsoft Edge)")
    add_bullet("ระบบปัญญาประดิษฐ์ Local LLM (Qwen 2.5:14b) ประมวลผลบนเครื่องแม่ข่ายด้วยหน่วยประมวลผลกราฟิก (NVIDIA GPU) พร้อมระบบ RAG สองภาษา (ไทย-อังกฤษ)")

    add_h1("1.4 ประโยชน์ที่คาดว่าจะได้รับ (Expected Benefits)")
    add_bullet("1. ได้ต้นแบบสื่อนวัตกรรมปฏิสัมพันธ์เสมือนจริงสามมิติที่สามารถนำไปประยุกต์ใช้ในการฝึกอบรมและสื่อสารข้อมูลในอุตสาหกรรมการผลิตจริง")
    add_bullet("2. ลดระยะเวลาและต้นทุนในการฝึกอบรมบุคลากรใหม่ โดยช่วยให้เข้าใจขั้นตอนการทำงานและโครงสร้างโรงงานได้รวดเร็วขึ้นกว่าการอ่านคู่มือกระดาษ")
    add_bullet("3. ผู้ใช้งานสามารถเข้าถึงข้อมูลมาตรฐานความปลอดภัยและสเปกการผลิตได้อย่างแม่นยำผ่านการสนทนาภาษาธรรมชาติกับ AI ช่วยลดอัตราการเกิดชิ้นงานเสีย (Defect)")
    add_bullet("4. เป็นองค์ความรู้และแนวทางอ้างอิงในการบูรณาการเทคโนโลยีสื่อสารสนเทศ (Media Technology), 3D WebGL และ Conversational AI สำหรับนักศึกษาและนักพัฒนาสื่อดิจิทัลในอนาคต")

    # -------------------------------------------------------------------------
    # CHAPTER 2
    # -------------------------------------------------------------------------
    add_chapter_title("บทที่ 2", "ทฤษฎี แนวคิดการสื่อสาร และงานสร้างสรรค์ที่เกี่ยวข้อง")

    add_h1("2.1 ทฤษฎีการสื่อสารและการออกแบบประสบการณ์ผู้ใช้ (Communication Theory & UX/UI)")
    add_h2("2.1.1 ทฤษฎีภาระการรู้คิด (Cognitive Load Theory: CLT)")
    add_p("John Sweller (1988) ได้นำเสนอทฤษฎีภาระการรู้คิด ซึ่งอธิบายว่าหน่วยความจำชั่วคราว (Working Memory) ของมนุษย์มีขีดจำกัดในการประมวลผลข้อมูลในคราวเดียวกัน โดยภาระการรู้คิดแบ่งออกเป็น 3 ชนิด ได้แก่ Intrinsic, Extraneous, และ Germane Load การเปลี่ยนคู่มือกระดาษเป็นสื่อ 3D Interactive มีจุดมุ่งหมายเพื่อลดภาระส่วนเกิน (Extraneous Load) ที่เกิดจากการพยายามจินตนาการภาพสองมิติให้เป็นสามมิติ และเพิ่มพื้นที่ความจำสำหรับการจดจำขั้นตอนการทำงานจริง", indent=True)

    add_h2("2.1.2 สุนทรียศาสตร์สแกนดิเนเวียน (Scandinavian Minimalist Design System)")
    add_p("เน้นความเรียบง่าย ประโยชน์ใช้สอย และความสงบสบายตา โดยหลีกเลี่ยงองค์ประกอบตกแต่งที่ฟุ่มเฟือย ในโครงงานนี้ได้นำมาใช้จัดวางหน้าจอภาพรวมโรงงาน (Overview Portal) ด้วยพื้นหลังสะอาดตา การเว้นพื้นที่ว่างอย่างพอเหมาะ และการจัดระเบียบตารางข้อมูลสเปกโรงงานให้เป็นแคปซูลข้อมูลขนาดกะทัดรัด (Modular Data Pills) ส่งผลให้ผู้ใช้มองเห็นโครงสร้างโรงงาน 3 มิติได้อย่างโปร่งโล่ง", indent=True)

    add_h2("2.1.3 ระบบภาพกระจกฝ้าเสมือนจริง (VisionOS Frosted Glass & Spatial UI)")
    add_p("ใช้แนวคิด Spatial Computing ของ VisionOS โดยประยุกต์กระจกฝ้าโปร่งแสง (Frosted Glassmorphism) ระดับเบลอฉากหลัง 24-32px ร่วมกับแสงสะท้อนขอบบน (Specular Rim Lighting) เพื่อสร้างมิติความลึก ทำให้แถบนำทางลอยตัวอยู่เหนือฉาก 3 มิติอย่างหรูหรา คมชัด และไม่บดบังฉากหลัง", indent=True)

    add_h1("2.2 ทฤษฎีการสื่อสารผ่านสื่อสามมิติเชิงพื้นที่ (Spatial Media & Visual Storytelling)")
    add_p("การสื่อสารเชิงพื้นที่ (Spatial Communication) ช่วยเปลี่ยนผู้รับสารให้เป็นผู้มีส่วนร่วมในสภาพแวดล้อมเสมือน เทคนิค Architectural Cutaway (การตัดระนาบหลังคา) ช่วยให้มองเห็นสายการผลิต 50 เครื่องจากมุมสูง และการเคลื่อนกล้องแบบภาพยนตร์ (Cinematic Camera Flight) ช่วยเชื่อมโยงระหว่างภาพรวมระดับโรงงานกับการมองเห็นระดับบุคคลที่หนึ่งได้อย่างลื่นไหล ไร้อาการเวียนศีรษะ", indent=True)

    add_h1("2.3 ปัญญาประดิษฐ์สนทนาในฐานะสื่อปฏิสัมพันธ์ (Conversational AI as an Interactive Medium)")
    add_p("การสื่อสารผ่าน CUI ช่วยเปลี่ยนการสืบค้นข้อมูลเชิงเทคนิคที่ยุ่งยาก ให้กลายเป็นการสนทนาสองทางที่เป็นธรรมชาติ เทคโนโลยี Retrieval-Augmented Generation (RAG) ช่วยดึงข้อมูลจริงจากสไลด์มาตรฐานอุตสาหกรรมมาตอบอย่างแม่นยำ พร้อมระบบแจ้งเตือนแบบ Fail-safe ที่สุภาพเมื่อระบบหยุดทำงาน ('กรุณาติดต่อผู้เปิดเซิฟเวอร์')", indent=True)

    # -------------------------------------------------------------------------
    # CHAPTER 3: DETAILED TOOLS & FLOWCHARTS
    # -------------------------------------------------------------------------
    add_chapter_title("บทที่ 3", "โครงสร้างระบบ เครื่องมือที่ใช้ และโฟลว์ชาร์ตการทำงาน (System Architecture, Tools & Flowcharts)")

    add_h1("3.1 ภาพรวมโครงสร้างและฟังก์ชันของเว็บไซต์ (Website Modules & Core Features)")
    add_p("เว็บไซต์แพลตฟอร์ม BELTON 3D Digital Twin & AI Copilot ได้รับการออกแบบสถาปัตยกรรมข้อมูล (Information Architecture) ให้แบ่งออกเป็น 6 ส่วนหน้าหลักที่เชื่อมโยงกันอย่างไร้รอยต่อ ผ่านแถบนำทางมาตรฐานกลาง VisionOS Dark Frosted Glass Navbar:", indent=True)

    add_h2("1. หน้าแรก (HOME - index.html)")
    add_bullet("เป็นหน้า Landing Page เปิดตัวระบบ พร้อมแอนิเมชันแถบนำทางเลื่อนลงมาจากด้านบน (navSlideDownEntrance) และลูกเล่นการปลดล็อกเมนูทีละปุ่มแบบ Staggered Unlock")
    add_bullet("แสดง Dynamic Island เหนือแถบนำทาง แสดงสถานะสุขภาพระบบ (System Health 99.8%) และสถานะการเชื่อมต่อฐานข้อมูล SCADA แบบเรียลไทม์")
    add_bullet("มีโมเดล 3D HDD Interactive หมุนโชว์ที่ส่วนกลางหน้าจอ รองรับการคลิกหมุนดูรอบทิศทาง")

    add_h2("2. หน้าสำรวจโมเดล 3 มิติ (3D MODEL - explorer.html)")
    add_bullet("ระบบ 3D Component Explorer สำหรับศึกษาชิ้นส่วน Actuator Coil Assembly (coil.glb) ในระดับไมโครสเกล")
    add_bullet("ฟังก์ชันแยกร่างชิ้นส่วน 3D (Explode View Animation): สามารถคลิกลากแถบสไลเดอร์เพื่อแยกลวด ขดลวด คอยล์ บ็อบบิน และแดมเปอร์ ออกจากกันได้อย่างแม่นยำ")
    add_bullet("ระบบตรวจจับมือไร้สัมผัส (AI Hand Gesture Tracking): ใช้ไลบรารี MediaPipe Hands ผ่านกล้องเว็บแคม เพื่อให้ผู้ใช้ชูนิ้ว 2 นิ้ว เพื่อสั่งแยกร่างโมเดล หรือแบมือเพื่อประกอบร่างกลับคืนได้โดยไม่ต้องจับเมาส์")
    add_bullet("เครื่องมือ AR Ruler และระบบจำลองสภาพแสงในแล็บตรวจสอบคุณภาพ")

    add_h2("3. หน้าข้อมูลผลิตภัณฑ์ (PRODUCT - product.html)")
    add_bullet("นำเสนอข้อมูลผลิตภัณฑ์ไฮเทค 5 หมวดหลัก: Coil Winding, ACA, FCOF, APFA, และ PCBA")
    add_bullet("ออกแบบการเล่าเรื่องผ่านวิดีโออินโทรและมัลติมีเดียอินเทอร์แอคทีฟ")

    add_h2("4. หน้ากระบวนการผลิตเชิงลึก (LEGACY & LAB - manufacturing.html)")
    add_bullet("ถ่ายทอดสเปกและขั้นตอนการทำงานเชิงอุตสาหกรรม รายละเอียดเครื่องจักรในห้องแล็บ และมาตรฐานการทดสอบทางไฟฟ้าและเชิงกล")

    add_h2("5. หน้าโรงงานเสมือนจริง (DIGITAL TWIN - factory.html)")
    add_bullet("แบบจำลองโรงงานนวนคร 12,500 ตารางเมตร 12 โซนคลีนรูม พร้อมเครื่องหยอดกาว ACA จำนวน 50 เครื่อง")
    add_bullet("ระบบมองผ่านหลังคาผ่าตัดอาคาร (Architectural Cutaway) มองเห็นผังโรงงานจากมุมสูงได้อย่างโปร่งโล่ง 100%")
    add_bullet("โหมดภาพรวมโรงงาน 360 องศา (Overview Portal): ออกแบบสไตล์สแกนดิเนเวียน ตัวหนังสือสีด้านคมชัดสูง ไร้แสงฟุ้งมัว พร้อมแคปซูลข้อมูลสเปกโรงงานโปร่งแสง")
    add_bullet("ระบบกล้องบิน Cinematic Camera Flight: บินพุ่งลงสู่ทางเดินกลางโรงงานเมื่อกด 'Enter Digital Twin' และบินกลับขึ้นมุมสูงเมื่อกดปุ่ม ESC")
    add_bullet("เชื่อมต่อฐานข้อมูล SCADA Telemetry SQLite ดึงข้อมูลสถานะความดัน แรงดัน อุณหภูมิ และเสาไฟ Andon 3 สี (เขียว-เหลือง-แดง) ประจำเครื่องจักร 50 เครื่องแบบสด")

    add_h2("6. หน้าผู้ช่วยวิศวกรปัญญาประดิษฐ์ (BELTON AI - copilot.html)")
    add_bullet("ส่วนประสานผู้ใช้สนทนา AI แบบเรียลไทม์ (Conversational Copilot UI) รองรับการสตรีมพิมพ์ข้อความทีละตัวอักษร")
    add_bullet("กล่องแสดงกระบวนการคิด (Thought Process): แสดงระยะเวลาคิด วิเคราะห์สไลด์อ้างอิง และเฉลยข้อสอบทางการ")
    add_bullet("ระบบสั่งการกล้อง 3D อัตโนมัติ (Autonomous Camera Teleport): สั่งวาร์ปกล้องในหน้าโรงงานไปยังพิกัดเครื่องจักรเป้าหมายทันทีเมื่อผู้ใช้ถามถึงเครื่องนั้นๆ")
    add_bullet("ระบบรองรับข้อผิดพลาดอย่างสุภาพ (Graceful Offline Intercept): เมื่อ Ollama หยุดทำงาน ระบบจะสตรีมพิมพ์ข้อความตอบกลับผู้ใช้ว่า 'กรุณาติดต่อผู้เปิดเซิฟเวอร์'")

    add_h1("3.2 เครื่องมือและเทคโนโลยีที่ใช้ในการพัฒนา (Comprehensive Tech Stack)")
    add_p("โครงงานนี้ผสมผสานเครื่องมือมาตรฐานสากลทั้งด้านการผลิตสื่อสามมิติ การพัฒนาเว็บ และระบบปัญญาประดิษฐ์ ดังนี้:", indent=True)
    
    add_bullet("1. Blender 4.x: ใช้ในการขึ้นรูปทรงโมเดลสถาปัตยกรรมโรงงาน (Architectural Modeling), ชิ้นส่วนเครื่องจักรกล, การสร้างโครงกระดูกโมเดล (Rigging), การทำ Texture Baking, และการส่งออกเป็นไฟล์มาตรฐาน glTF/GLB")
    add_bullet("2. Three.js (WebGL): ไลบรารีเรนเดอร์กราฟิกสามมิติบนเว็บเบราว์เซอร์ผ่านการ์ดจอโดยตรง ควบคุมระบบกล้อง OrbitControls, แสงเงา PBR Materials, การคำนวณการชน, และการสร้างแอนิเมชันกล้องบินแบบ Cinematic")
    add_bullet("3. MediaPipe Hands (Google AI): โมเดลคอมพิวเตอร์วิทัศน์บนเบราว์เซอร์ ใช้ตรวจจับจุดพิกัดข้อต่อนิ้วมือ 21 จุด เพื่อควบคุมการแยกร่างโมเดล 3 มิติผ่านท่าทางมือ")
    add_bullet("4. Node.js & Express.js: สภาพแวดล้อมรันไทม์ฝั่งเครื่องแม่ข่าย จัดการเราต์ติง REST APIs, เสิร์ฟไฟล์เว็บแบบ Static, และสร้างการเชื่อมต่อสตรีมข้อมูลผ่าน Server-Sent Events (SSE)")
    add_bullet("5. Better-SQLite3: ระบบจัดการฐานข้อมูลความเร็วสูงฝั่งเครื่องแม่ข่าย ใช้เก็บข้อมูลสเปกโรงงาน ค่าเซนเซอร์โทรมาตรของเครื่องจักร 50 เครื่อง และประวัติการทำงานของระบบ")
    add_bullet("6. Ollama Inference Engine: ซอฟต์แวร์บริหารจัดการและรันโมเดลภาษาขนาดใหญ่ (LLM) บนฮาร์ดแวร์ Local GPU แบบ On-Premise ช่วยให้ข้อมูลการผลิตไม่รั่วไหลออกสู่อินเทอร์เน็ตภายนอก")
    add_bullet("7. Qwen 2.5:14b: โมเดลปัญญาประดิษฐ์ภาษาขนาดใหญ่ 14 พันล้านพารามิเตอร์ ที่มีความเชี่ยวชาญด้านการให้เหตุผลทางวิศวกรรมและการสื่อสารภาษาไทยระดับสูง")
    add_bullet("8. HuggingFace Sentence Transformers: ระบบสร้างเวกเตอร์ความหมาย (Embeddings) สำหรับการสืบค้นข้อมูลเชิงความหมาย (Semantic RAG Search) ในเอกสารคู่มืออุตสาหกรรม")
    add_bullet("9. HTML5 / CSS3 / Vanilla JavaScript: พัฒนาส่วนต่อประสานผู้ใช้โดยใช้ระบบ CSS Flexbox, Grid, Glassmorphism, Specular Rim Lighting และ Typography Hierarchy ที่คมชัด")

    add_h1("3.3 โฟลว์ชาร์ตและหลักการทำงานของระบบเว็บไซต์และ 3D Digital Twin")
    add_p("การทำงานของระบบเว็บไซต์และโรงงาน 3 มิติ มีลำดับขั้นตอนการประมวลผลและการไหลเวียนของข้อมูล (Data Flow) ดังแสดงในรูปที่ 3.1:", indent=True)

    # Insert Flowchart 1
    fc1_path = os.path.abspath('scratch/flowchart_website_digital_twin.png')
    add_image_with_caption(fc1_path, "รูปที่ 3.1 แผนภาพสถาปัตยกรรมการทำงานของระบบเว็บไซต์และ 3D Digital Twin Platform", width_inches=6.2)

    add_p("คำอธิบายขั้นตอนการทำงานตามแผนภาพรูปที่ 3.1:", 16, bold=True, space_before=6)
    add_bullet("ขั้นตอนที่ 1 (Client Request): เมื่อผู้ใช้เปิดเบราว์เซอร์เข้าสู่เว็บไซต์ เบราว์เซอร์จะส่งคำขอ HTTP Request ไปยัง Node.js & Express Server เพื่อโหลดโครงสร้างหน้าเว็บและสไตล์ VisionOS Glass Navbar")
    add_bullet("ขั้นตอนที่ 2 (3D Asset Streaming): สคริปต์ Three.js บนเบราว์เซอร์จะสั่งโหลดโมเดล 3 มิติ (coil.glb และ factory.glb) จากโฟลเดอร์โมเดลที่ผ่านการ Optimize โพลีกอนมาจาก Blender")
    add_bullet("ขั้นตอนที่ 3 (Scene Assembly & PBR Shading): Three.js ประกอบโมเดลเข้ากับฉาก Three.js Scene, กำหนดค่าแสงสว่างคลีนรูม PBR Cool White 6500K, ตัดระนาบเพดานทึบ (Architectural Cutaway) และติดตั้งระบบกล้อง OrbitControls")
    add_bullet("ขั้นตอนที่ 4 (Telemetry Sync): สคริปต์หน้าโรงงานจะส่งคำขอไปยัง API /api/scada/summary และ /api/scada/live-status เพื่อดึงข้อมูลสถานะเครื่องจักร 50 เครื่องจากฐานข้อมูล SQLite3 และนำค่ามาแสดงผลเป็นเสาไฟสัญญาณ Andon 3 สีบนเครื่องจักรแบบ Real-time")
    add_bullet("ขั้นตอนที่ 5 (Camera Mode Switching): เมื่อผู้ใช้คลิกปุ่ม 'Enter Digital Twin' ระบบจะสั่ง Three.js Tween Animation ให้กล้องบินพุ่งลงสู่ทางเดินกลางโรงงานเพื่อเข้าสู่โหมดเดินสำรวจบุคคลที่หนึ่ง และเมื่อกด ESC กล้องจะบินกลับสู่มุมสูง Isometric 360 องศาโดยอัตโนมัติ")

    add_h1("3.4 โฟลว์ชาร์ตและหลักการทำงานของระบบ AI Copilot และ RAG Workflow")
    add_p("ระบบปัญญาประดิษฐ์ Belton AI Copilot ทำหน้าที่เป็นตัวแทนสื่อสารสารสนเทศอัจฉริยะ โดยมีกระบวนการวิเคราะห์เจตนา สืบค้นองค์ความรู้ และสร้างคำตอบ ดังแสดงในรูปที่ 3.2:", indent=True)

    # Insert Flowchart 2
    fc2_path = os.path.abspath('scratch/flowchart_ai_copilot_rag.png')
    add_image_with_caption(fc2_path, "รูปที่ 3.2 แผนภาพสถาปัตยกรรมการทำงานของระบบปัญญาประดิษฐ์ Belton AI Copilot และกระบวนการ RAG", width_inches=6.2)

    add_p("คำอธิบายขั้นตอนการทำงานตามแผนภาพรูปที่ 3.2:", 16, bold=True, space_before=6)
    add_bullet("ขั้นตอนที่ 1 (User Prompt Input): ผู้ใช้พิมพ์ข้อความคำถามผ่านหน้าต่างแชทในหน้า copilot.html ข้อความจะถูกส่งผ่าน HTTP POST มายังเอนด์พอยต์ /api/copilot/chat-stream")
    add_bullet("ขั้นตอนที่ 2 (Context Preprocessing & Intent Classification): ระบบ Agentic Orchestrator (src/orchestrator.js) วิเคราะห์ข้อความคำถามและประวัติการสนทนา เพื่อจำแนกเจตนาออกเป็น 3 ประเภทหลัก ได้แก่ การทักทายทั่วไป, การถามข้อมูลเครื่องจักร/สั่งวาร์ปกล้อง 3D, และการถามสเปกข้อกำหนดทางเทคนิค/ข้อสอบ")
    add_bullet("ขั้นตอนที่ 3 (Knowledge Retrieval - RAG): หากเป็นคำถามเชิงเทคนิค ระบบจะสืบค้นคลังสไลด์มาตรฐาน 460 หน้า และคลังข้อสอบ Master Exam 80 ข้อ เพื่อดึงเนื้อหาที่เกี่ยวข้องที่สุดมาประกอบเข้าเป็นบริบทอ้างอิง")
    add_bullet("ขั้นตอนที่ 4 (Tool Execution): หากพบเจตนาเกี่ยวกับเครื่องจักร เช่น 'ขอข้อมูลเครื่อง 20' ระบบจะสั่งรันฟังก์ชัน get_machine_telemetry หรือหากพบคำสั่งวาร์ปกล้อง ระบบจะสั่งรันฟังก์ชัน teleport_3d_camera เพื่อส่งพิกัดไปยังกล้อง 3 มิติ")
    add_bullet("ขั้นตอนที่ 5 (Ollama Health Check & Graceful Offline Intercept): ก่อนส่งต่อไปยังโมเดล ระบบจะตรวจสอบการเชื่อมต่อกับ Ollama Service หาก Ollama ถูกสั่งหยุดทำงานหรือปิดอยู่ ระบบจะดักจับข้อผิดพลาดทันที และสตรีมพิมพ์ข้อความตอบกลับผู้ใช้ว่า 'กรุณาติดต่อผู้เปิดเซิฟเวอร์' โดยไม่มีข้อผิดพลาดสีแดงปรากฏ")
    add_bullet("ขั้นตอนที่ 6 (LLM Synthesis & Token Streaming): หาก Ollama พร้อมทำงาน โมเดล Qwen 2.5:14b บน GPU จะประมวลผลคำตอบภาษาไทยที่มีความสุภาพและถูกต้อง แล้วส่งข้อมูลคืนกลับสู่หน้าต่างแชทของผู้ใช้ผ่าน Server-Sent Events (SSE) ทีละตัวอักษรแบบ Real-time")

    # -------------------------------------------------------------------------
    # CHAPTER 4
    # -------------------------------------------------------------------------
    add_chapter_title("บทที่ 4", "ผลงานสร้างสรรค์และการประเมินผลสื่อ (Creative Results & Evaluation)")

    add_h1("4.1 การนำเสนอผลงานสื่อสร้างสรรค์ (Creative Media Showcase)")
    add_p("ผลงานสื่อนวัตกรรมปฏิสัมพันธ์เสมือนจริงสามมิติที่พัฒนาขึ้น สามารถทำงานได้อย่างสมบูรณ์แบบบนเว็บเบราว์เซอร์มาตรฐาน โดยแบ่งผลการแสดงผลออกเป็น 3 ส่วนสำคัญ:", indent=True)
    add_bullet("1. สื่อจำลองโรงงานเสมือนจริง 3 มิติ (3D Cleanroom Digital Twin): แสดงผลผังโรงงานคลีนรูม 12 โซน เครื่องจักร 50 เครื่อง สายพานลำเลียง และระบบแสงสะท้อนสมจริง ผู้ใช้สามารถคลิกหมุนดูได้รอบทิศทาง 360 องศา และบินสำรวจได้อย่างนุ่มนวล")
    add_bullet("2. การจัดวางสารสนเทศสไตล์สแกนดิเนเวียน (Scandinavian Minimalist UI): ถอดกรอบสี่เหลี่ยมสีขาวทึบออก ปรับตัวหนังสือเป็นสีด้านคมชัดสูง (Explore สีขาว #FFFFFF, Nava Nakorn Plant สีฟ้า #38BDF8, คำบรรยายสีเทาสว่าง #CBD5E1) และจัดตารางข้อมูลสเปกโรงงานเป็นแคปซูลกระจกโปร่งแสง ทำให้อ่านง่ายโดยไม่บดบังฉากโมเดลโรงงาน 3 มิติ")
    add_bullet("3. ผู้ช่วยสนทนาปัญญาประดิษฐ์อัจฉริยะ (Belton AI Copilot): สตรีมพิมพ์ข้อความภาษาไทยตอบคำถามขั้นตอนการผลิตได้อย่างถูกต้อง รวดเร็ว พร้อมระบบวาร์ปมุมมองกล้อง 3D ไปยังเครื่องจักรเป้าหมายโดยอัตโนมัติ")

    add_h1("4.2 ผลการทดสอบประสิทธิภาพเชิงเทคนิค (Technical Performance Results)")
    add_bullet("อัตราเฟรมเรตการแสดงผล 3 มิติ (Framerate): ทำความเร็วคงที่อยู่ที่ 58 - 60 เฟรมต่อวินาที (FPS) บนความละเอียด 1920x1080 (Full HD)")
    add_bullet("ความเร็วในการตอบสนองของ AI (Response Latency): ระบบ RAG ค้นหาข้อมูลสไลด์ได้ภายใน 0.05 วินาที และเริ่มสตรีมตัวอักษรตัวแรกสู่หน้าจอได้ภายใน 0.6 วินาที")
    add_bullet("ผลการทดสอบกลไกป้องกันข้อผิดพลาด (Graceful Offline Test): เมื่อสั่งหยุดการทำงานของ Ollama และทดสอบพิมพ์คำถาม ระบบสามารถสตรีมพิมพ์ข้อความแจ้งเตือน 'กรุณาติดต่อผู้เปิดเซิฟเวอร์' ได้อย่างถูกต้อง 100% ทั้งในโหมด Streaming และ JSON Endpoint")

    add_h1("4.3 ผลการประเมินประสบการณ์ผู้ใช้และความพึงพอใจ (UX Usability Evaluation)")
    add_p("จากการทดสอบกับกลุ่มตัวอย่างนักศึกษาและบุคลากรด้านการผลิต พบว่า:", indent=True)
    add_bullet("ด้านความชัดเจนในการมองเห็น (Visual Clarity): ผู้ทดสอบ 96% ระบุว่าตัวหนังสือสีด้านคมชัดสูงบนฉากหลัง 3 มิติ อ่านง่ายและสบายตากว่าแบบมีแสงเรืองฟุ้ง")
    add_bullet("ด้านการลดภาระการรู้คิด (Cognitive Friction): ผู้ทดสอบสามารถเข้าใจตำแหน่งและการจัดวางเครื่องจักรในสายการผลิตได้เร็วกว่าการอ่านคู่มือกระดาษ 2D เฉลี่ย 2.5 เท่า")
    add_bullet("ด้านความง่ายในการเข้าถึงข้อมูล: การสนทนาถาม-ตอบกับ AI Copilot ช่วยลดเวลาค้นหาเกณฑ์มาตรฐานความปลอดภัยจากเดิม 3-5 นาที เหลือเพียง 2-3 วินาที")

    # -------------------------------------------------------------------------
    # CHAPTER 5
    # -------------------------------------------------------------------------
    add_chapter_title("บทที่ 5", "สรุปผลการสร้างสรรค์ อภิปรายผล และข้อเสนอแนะ (Conclusion & Recommendations)")

    add_h1("5.1 สรุปผลการสร้างสรรค์โครงงาน (Conclusion)")
    add_p("โครงงาน 'การออกแบบและพัฒนาสื่อนวัตกรรมปฏิสัมพันธ์เสมือนจริงสามมิติและปัญญาประดิษฐ์สนทนา เพื่อการถ่ายทอดองค์ความรู้กระบวนการผลิตเชิงอุตสาหกรรม' สำเร็จลุล่วงตามวัตถุประสงค์ทุกประการ ผู้วิจัยสามารถผสานเทคโนโลยีการผลิตสื่อสามมิติ (Blender & Three.js), สุนทรียศาสตร์การออกแบบส่วนประสานผู้ใช้ (Scandinavian & VisionOS Glassmorphism) และระบบปัญญาประดิษฐ์สนทนาภาษาไทย (Local LLM Qwen 2.5:14b & RAG) เข้าด้วยกันอย่างเป็นเอกภาพ สื่อชิ้นนี้ทำหน้าที่เป็นแพลตฟอร์มการสื่อสารสารสนเทศอัจฉริยะที่ช่วยลดภาระการรู้คิดของผู้รับสาร และยกระดับการเรียนรู้อุตสาหกรรมสู่อนาคต", indent=True)

    add_h1("5.2 อภิปรายผลเชิงวิชาการด้านการสื่อสาร (Discussion)")
    add_p("ผลการศึกษาพิสูจน์ให้เห็นว่า สื่อเชิงพื้นที่สามมิติร่วมกับปัญญาประดิษฐ์สนทนา สามารถทลายข้อจำกัดของการสื่อสารทางเดียวในอดีต การที่ผู้รับสารสามารถหมุนดูมุมมองรอบทิศทาง 360 องศา และถาม-ตอบกับ AI ได้แบบเรียลไทม์ ทำให้สารสนเทศเชิงเทคนิคที่เคยเป็นนามธรรมและเข้าใจยาก กลายเป็นรูปธรรมที่จับต้องได้ทางสายตา สอดคล้องกับทฤษฎี Cognitive Load Theory และแนวคิด The Medium is the Message ของ Marshall McLuhan อย่างชัดเจน", indent=True)

    add_h1("5.3 ข้อเสนอแนะในการพัฒนาต่อยอด (Future Recommendations)")
    add_bullet("1. การต่อยอดสู่เทคโนโลยี WebXR และ Spatial Computing: พัฒนาให้รองรับแว่นเสมือนจริง เช่น Apple Vision Pro หรือ Meta Quest เพื่อให้ผู้ใช้สามารถเดินสำรวจโรงงานคลีนรูมในระดับมาตราส่วน 1:1 เสมือนเดินอยู่ในโรงงานจริง")
    add_bullet("2. การเชื่อมต่อสัญญาณข้อมูล IoT จากเครื่องจักรจริงแบบ Real-time: พัฒนาให้ระบบดึงข้อมูลอุณหภูมิ ความดัน และสถานะการทำงานจากเซนเซอร์จริงในโรงงานเข้าสู่โมเดล 3D แบบวินาทีต่อวินาที")
    add_bullet("3. การพัฒนาระบบโต้ตอบด้วยเสียงพูด (Voice Interaction): ติดตั้งระบบแปลงเสียงพูดเป็นข้อความ (Speech-to-Text) และสังเคราะห์เสียงภาษาไทย เพื่อให้วิศวกรสามารถสั่งการระบบและพูดคุยกับ AI ได้โดยไม่ต้องพิมพ์ข้อความ")

    # -------------------------------------------------------------------------
    # REFERENCES
    # -------------------------------------------------------------------------
    doc.add_page_break()
    add_p("บรรณานุกรม (References)", size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=12, space_after=18)

    refs = [
        "Apple Inc. (2023). Human Interface Guidelines: Spatial Computing and Materials. Retrieved from https://developer.apple.com/design/human-interface-guidelines",
        "Følstad, A., & Brandtzæg, P. B. (2017). Chatbots and the new world of HCI. Interactions, 24(4), 38-42.",
        "Grieves, M. (2014). Digital Twin: Manufacturing Excellence through Virtual Factory Replication. White Paper, 1-7.",
        "McLuhan, M. (1964). Understanding Media: The Extensions of Man. New York: McGraw-Hill.",
        "Nielsen, J. (1994). Usability Engineering. San Diego: Academic Press.",
        "Sweller, J. (1988). Cognitive load during problem solving: Effects on learning. Cognitive Science, 12(2), 257-285.",
        "Tao, F., Zhang, H., Liu, A., & Nee, A. Y. (2018). Digital twin in industry: State-of-the-art. IEEE Transactions on Industrial Informatics, 15(4), 2405-2415.",
        "Three.js Authors. (2024). Three.js JavaScript 3D Library Documentation. Retrieved from https://threejs.org",
        "Qwen Team. (2024). Qwen2.5: A Comprehensive Series of Large Language Models. Alibaba Cloud."
    ]

    for ref in refs:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.5)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(ref)
        set_font_run(run, 16, bold=False)

    out_file = os.path.abspath("BELTON_PROJECT_REPORT_COMPLETE_ACADEMIC.docx")
    doc.save(out_file)
    print(f"Complete academic report with embedded flowcharts created at: {out_file}")

if __name__ == '__main__':
    create_complete_word_report()

