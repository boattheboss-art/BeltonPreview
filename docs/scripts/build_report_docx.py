import shutil
import os
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

src_md = r'c:\Users\BOAT\Videos\เลขา\belton_live_preview\PROJECT_REPORT_5_CHAPTERS.md'
dst_md = r'C:\Users\BOAT\Downloads\PROJECT_REPORT_5_CHAPTERS.md'
dst_docx1 = r'C:\Users\BOAT\Downloads\PROJECT_REPORT_5_CHAPTERS.docx'
dst_docx2 = r'C:\Users\BOAT\Downloads\PROJECT_REPORT_5_CHAPTERS_REVISED.docx'

# 1. Copy Markdown to Downloads
shutil.copyfile(src_md, dst_md)
print('Copied markdown to Downloads.')

# 2. Build Document
doc = Document()
for section in doc.sections:
    section.top_margin = Inches(1.5)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.5)
    section.right_margin = Inches(1.0)

# Disable spelling & grammar errors across the entire document
try:
    settings_elm = doc.settings.element
    settings_elm.append(parse_xml(f'<w:hideSpellingErrors {nsdecls("w")}/>'))
    settings_elm.append(parse_xml(f'<w:hideGrammaticalErrors {nsdecls("w")}/>'))
except Exception as e:
    print('Warning settings:', e)

# Base style setup with Thai proofing language and noProof
style_normal = doc.styles['Normal']
style_normal.font.name = 'TH Sarabun New'
style_normal.font.size = Pt(16)
style_normal.font.color.rgb = RGBColor(0, 0, 0)
try:
    rPr = style_normal.element.get_or_add_rPr()
    rPr.append(parse_xml(f'<w:noProof {nsdecls("w")}/>'))
    rPr.append(parse_xml(f'<w:lang {nsdecls("w")} w:val="th-TH" w:eastAsia="th-TH" w:bidi="th-TH"/>'))
except Exception as e:
    print('Warning style rPr:', e)

def apply_thai_no_proof(run):
    try:
        r_elm = run._r.get_or_add_rPr()
        r_elm.append(parse_xml(f'<w:noProof {nsdecls("w")}/>'))
        r_elm.append(parse_xml(f'<w:lang {nsdecls("w")} w:val="th-TH" w:bidi="th-TH"/>'))
    except Exception:
        pass

with open(src_md, 'r', encoding='utf-8') as f:
    text_content = f.read()

# Separate cover page from the rest of the document
parts = text_content.split('<div style="page-break-after: always;"></div>')
cover_raw = parts[0]
body_raw = '<div style="page-break-after: always;"></div>'.join(parts[1:])

# =========================================================================
# 1. RENDER COVER PAGE (STRICTLY CENTER-ALIGNED WITH ACADEMIC PROPORTIONS)
# =========================================================================

# Title Block at Top
p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_title.paragraph_format.space_before = Pt(40)
p_title.paragraph_format.space_after = Pt(12)
p_title.paragraph_format.line_spacing = 1.25
r_title = p_title.add_run('การออกแบบและพัฒนาสื่อนวัตกรรมปฏิสัมพันธ์เสมือนจริงสามมิติ\nและปัญญาประดิษฐ์สนทนา เพื่อการถ่ายทอดองค์ความรู้\nกระบวนการผลิตเชิงอุตสาหกรรม')
r_title.bold = True
r_title.font.name = 'TH Sarabun New'
r_title.font.size = Pt(20)
apply_thai_no_proof(r_title)

# Author Block (Centered, Middle Page)
p_author = doc.add_paragraph()
p_author.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_author.paragraph_format.space_before = Pt(85)
p_author.paragraph_format.space_after = Pt(4)
p_author.paragraph_format.line_spacing = 1.2
r_by = p_author.add_run('โดย\n')
r_by.font.name = 'TH Sarabun New'
r_by.font.size = Pt(16)
apply_thai_no_proof(r_by)
r_name = p_author.add_run('นาย [ระบุชื่อ-นามสกุล ผู้จัดทำ]\n')
r_name.bold = True
r_name.font.name = 'TH Sarabun New'
r_name.font.size = Pt(18)
apply_thai_no_proof(r_name)
r_id = p_author.add_run('รหัสนักศึกษา [ระบุรหัสนักศึกษา]')
r_id.font.name = 'TH Sarabun New'
r_id.font.size = Pt(16)
apply_thai_no_proof(r_id)

# Advisor Block (Centered)
p_adv = doc.add_paragraph()
p_adv.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_adv.paragraph_format.space_before = Pt(50)
p_adv.paragraph_format.space_after = Pt(4)
p_adv.paragraph_format.line_spacing = 1.2
r_adv_label = p_adv.add_run('อาจารย์ที่ปรึกษาโครงงาน\n')
r_adv_label.font.name = 'TH Sarabun New'
r_adv_label.font.size = Pt(16)
apply_thai_no_proof(r_adv_label)
r_adv_val = p_adv.add_run('[ระบุชื่อ-นามสกุล อาจารย์ที่ปรึกษา]')
r_adv_val.bold = True
r_adv_val.font.name = 'TH Sarabun New'
r_adv_val.font.size = Pt(18)
apply_thai_no_proof(r_adv_val)

# University Degree Block (Centered, Bottom Page)
p_uni = doc.add_paragraph()
p_uni.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_uni.paragraph_format.space_before = Pt(85)
p_uni.paragraph_format.space_after = Pt(4)
p_uni.paragraph_format.line_spacing = 1.2
r_u1 = p_uni.add_run('รายงานโครงงานนี้เป็นส่วนหนึ่งของการศึกษาตามหลักสูตรปริญญาบัณฑิต\n')
r_u1.bold = True
r_u1.font.name = 'TH Sarabun New'
r_u1.font.size = Pt(16)
apply_thai_no_proof(r_u1)
r_u2 = p_uni.add_run('สาขาวิชา [ระบุสาขาวิชา] คณะ [ระบุคณะ]\n[ระบุชื่อมหาวิทยาลัย / สถาบันการศึกษา]\nปีการศึกษา 2567')
r_u2.font.name = 'TH Sarabun New'
r_u2.font.size = Pt(16)
apply_thai_no_proof(r_u2)

# PAGE BREAK AFTER COVER PAGE
doc.add_page_break()

# =========================================================================
# 2. RENDER BODY CONTENT WITH PROPER HEADINGS AND PAGE BREAKS
# =========================================================================
lines = body_raw.split('\n')
in_code = False
code_buffer = []

def add_styled_runs(paragraph, text, base_size=16, base_bold=False):
    tokens = re.split(r'(\*\*.*?\*\*)', text)
    for token in tokens:
        if token.startswith('**') and token.endswith('**'):
            run = paragraph.add_run(token[2:-2])
            run.font.name = 'TH Sarabun New'
            run.font.size = Pt(base_size)
            run.bold = True
        else:
            run = paragraph.add_run(token)
            run.font.name = 'TH Sarabun New'
            run.font.size = Pt(base_size)
            run.bold = base_bold
        apply_thai_no_proof(run)

for line in lines:
    raw = line.rstrip('\r\n')
    s_raw = raw.strip()

    # Explicit page break
    if '<div style="page-break-after: always;"></div>' in raw:
        doc.add_page_break()
        continue

    # Skip HTML wrapper tags
    if s_raw in ['<div align="center">', '</div>', '<br>', '<br/>', '<hr>']:
        continue

    # Code block handling
    if s_raw.startswith('```'):
        if in_code:
            in_code = False
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run('\n'.join(code_buffer))
            r.font.name = 'Consolas'
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(60, 60, 60)
            apply_thai_no_proof(r)
            code_buffer = []
        else:
            in_code = True
            code_buffer = []
        continue

    if in_code:
        code_buffer.append(raw)
        continue

    # Skip markdown hr
    if s_raw == '---':
        continue

    # Chapter / Major Heading (# )
    if raw.startswith('# '):
        clean_title = raw[2:].replace('<br>', '\n').strip()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(clean_title)
        r.bold = True
        r.font.name = 'TH Sarabun New'
        r.font.size = Pt(18)
        apply_thai_no_proof(r)

    # Section Heading (## )
    elif raw.startswith('## '):
        clean_title = raw[3:].strip()
        p = doc.add_paragraph()
        if any(w in clean_title for w in ['สารบัญ', 'บรรณานุกรม', 'ภาคผนวก']):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(clean_title)
        r.bold = True
        r.font.name = 'TH Sarabun New'
        r.font.size = Pt(16)
        apply_thai_no_proof(r)

    # Subsection Heading (### )
    elif raw.startswith('### '):
        clean_title = raw[4:].strip()
        p = doc.add_paragraph()
        if 'สารบัญ' in clean_title:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(clean_title)
        r.bold = True
        r.font.name = 'TH Sarabun New'
        r.font.size = Pt(15)
        apply_thai_no_proof(r)

    # Table rows
    elif s_raw.startswith('|') and '|' in s_raw[1:]:
        if '---' in s_raw:
            continue
        cells = [c.strip() for c in s_raw.strip('|').split('|')]
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(' • ' + '  |  '.join(cells))
        r.font.name = 'TH Sarabun New'
        r.font.size = Pt(11)
        apply_thai_no_proof(r)

    # Unordered Lists
    elif s_raw.startswith('- ') or s_raw.startswith('* '):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        add_styled_runs(p, s_raw[2:].strip())

    # Numbered Lists
    elif re.match(r'^\d+\.\s', s_raw):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        add_styled_runs(p, s_raw)

    elif s_raw == '':
        continue

    # Standard Paragraphs
    else:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if not s_raw.startswith('**') and not s_raw.startswith('ตารางที่') and not s_raw.startswith('ภาพที่'):
            p.paragraph_format.first_line_indent = Inches(0.5)
        elif s_raw.startswith('ตารางที่') or s_raw.startswith('ภาพที่'):
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        add_styled_runs(p, raw)

# Save to both target DOCX files
saved_count = 0
for target_path in [dst_docx1, dst_docx2]:
    try:
        doc.save(target_path)
        print('Successfully saved:', target_path)
        saved_count += 1
    except Exception as e:
        print('Could not save to', target_path, ':', e)

print(f'Done! Successfully saved {saved_count} Word document(s) with noProof embedded.')
